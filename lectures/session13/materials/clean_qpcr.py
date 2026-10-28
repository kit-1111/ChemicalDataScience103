"""Generated solution: Clean qpcr_exp14.xlsx into a tidy, analysis-ready CSV.

Every judgement call is recorded in the LOG list and written to cleaning_log.txt,
so the cleaning is documented and rerunnable. Raw file is never modified.
"""
import re
from datetime import date, datetime

import numpy as np
import pandas as pd
from openpyxl import load_workbook

SRC = "qpcr_exp14.xlsx"
OUT = "qpcr_exp14_tidy.csv"
LOG_OUT = "cleaning_log.txt"
LOG = []

# ---------------------------------------------------------------- 1. read
# header=None: the sheet has titles, merged cells and a two-row header, so we
# locate the table ourselves instead of trusting pandas to guess.
raw = pd.read_excel(SRC, sheet_name="Results", header=None)
FIRST, LAST = 5, 19          # 0-based rows of the data block (Excel rows 6-20)
body = raw.iloc[FIRST:LAST + 1, :10].copy()
body.columns = ["sample_raw", "g1", "g2", "GAPDH", "ACTB",
                "conc_raw", "vol_raw", "date_raw", "comment", "run_id"]
body = body.reset_index(drop=True)

# Gene names: Excel turned SEPT2 / MARCH1 into dates. This cannot be reversed
# automatically (the original text is gone), so we restore them from a lookup.
# Newer HGNC symbols are SEPTIN2 / MARCHF1; we keep the names used in the lab.
hdr = raw.iloc[4, 1:3].tolist()
assert all(isinstance(h, (datetime, pd.Timestamp)) for h in hdr), "header layout changed"
body = body.rename(columns={"g1": "SEPT2", "g2": "MARCH1"})
LOG.append(f"Header cells {[h.strftime('%d-%b') for h in hdr]} were dates; restored to SEPT2, MARCH1.")

# ---------------------------------------------------------------- 2. formatting-only information
# Colour carries meaning in this sheet, and pandas cannot see it. Read it with openpyxl.
ws = load_workbook(SRC)["Results"]
def fill_of(excel_row):
    c = ws.cell(excel_row, 1).fill
    return (c.fgColor.rgb or "")[-6:] if c.fill_type else ""
fills = [fill_of(r) for r in range(FIRST + 1, LAST + 2)]   # Excel rows 6-20
body["excluded"] = [f == "FFC7CE" for f in fills]       # red: "don't use!!" (note K5)
body["repeat_if_time"] = [f == "FFFF00" for f in fills]   # yellow (note K7)
LOG.append("Red rows -> excluded=True; yellow cells -> repeat_if_time=True (meaning taken from margin notes K5, K7). "
           "Reason for exclusion is NOT documented in the source: ask the experimenter.")

# ---------------------------------------------------------------- 3. sample labels
def parse_sample(s):
    t = str(s).strip().lower().replace("µ", "u")
    nums = re.findall(r"\d+", t)
    if re.match(r"(ctrl|control)", t):
        return pd.Series({"treatment": "control", "dose_uM": 0.0, "replicate": int(nums[-1])})
    m = re.match(r"drug\s*([ab])", t)
    dose = re.search(r"(\d+)\s*um", t)
    return pd.Series({"treatment": f"drug_{m.group(1)}", "dose_uM": float(dose.group(1)),
                      "replicate": int(nums[-1])})   # replicate = last integer in the label

body = pd.concat([body, body["sample_raw"].apply(parse_sample)], axis=1)
body["sample_id"] = (body.treatment + "_" + body.dose_uM.astype(int).astype(str) + "uM_r"
                     + body.replicate.astype(str))
assert body.sample_id.is_unique, "sample parsing produced duplicate IDs"

# ---------------------------------------------------------------- 4. Ct values
def parse_ct(x):
    """Return (value, flag). Non-numeric codes become NaN plus an explicit flag."""
    if pd.isna(x):
        return np.nan, "missing"
    if isinstance(x, (int, float)):
        return float(x), ""
    s = str(x).strip().lower()
    if s == "undetermined":
        return np.nan, "undetermined"
    if s == "n.d.":
        return np.nan, "not_detected"
    flag = "uncertain" if "(?)" in s else ""
    s = s.replace("(?)", "").strip().replace(",", ".")      # decimal comma -> point
    return float(s), flag

genes = ["SEPT2", "MARCH1", "GAPDH", "ACTB"]
for g in genes:
    parsed = body[g].apply(parse_ct)
    body[f"{g}_ct"] = [p[0] for p in parsed]
    body[f"{g}_flag"] = [p[1] for p in parsed]
LOG.append("Ct: decimal commas converted; '(?)' kept as flag 'uncertain'; 'Undetermined', 'n.d.' and blank "
           "kept as NaN with distinct flags (they are not the same thing). Nothing was imputed.")

# ---------------------------------------------------------------- 5. units
def conc_mg_per_ml(x):
    m = re.match(r"\s*([\d.,]+)\s*(mg|g|µg|ug)\s*/\s*ml", str(x), flags=re.I)
    val = float(m.group(1).replace(",", "."))
    return val * {"mg": 1, "g": 1000, "µg": 1e-3, "ug": 1e-3}[m.group(2).lower()]

def vol_ul(x):
    s = str(x).lower().replace("µ", "u")
    m = re.match(r"\s*([\d.,]+)\s*(ul|ml)?", s)
    val = float(m.group(1).replace(",", "."))
    return val * (1000 if m.group(2) == "ml" else 1)         # bare number assumed to be µL

body["lysate_mg_per_ml"] = body.conc_raw.apply(conc_mg_per_ml)
body["volume_ul"] = body.vol_raw.apply(vol_ul)
LOG.append("Concentration converted to mg/mL (g/mL x1000, ug/mL /1000). Volume to uL; bare numbers assumed to be uL.")

# ---------------------------------------------------------------- 6. dates
def parse_date(x):
    if isinstance(x, (datetime, pd.Timestamp, date)):
        return pd.Timestamp(x).date()
    s = str(x).strip()
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return datetime.strptime(s + " 2025", "%B %d %Y").date()   # "March 13": year from sheet header

body["measured_date"] = body.date_raw.apply(parse_date)
LOG.append("Dates: dotted = day.month; slashed '3/12/2025' read as month/day (= 12 March, the experiment date, "
           "while day/month would give 3 December). 'March 13' got year 2025 from the sheet header. CHECK with experimenter.")

# ---------------------------------------------------------------- 7. wide -> long (tidy)
meta = ["sample_id", "treatment", "dose_uM", "replicate", "lysate_mg_per_ml", "volume_ul",
        "measured_date", "run_id", "excluded", "repeat_if_time", "comment"]
long = pd.concat(
    [body[meta].assign(gene=g, ct=body[f"{g}_ct"], ct_flag=body[f"{g}_flag"]) for g in genes],
    ignore_index=True,
)
long["run_id"] = long["run_id"].astype("string")
long["comment"] = long["comment"].fillna("")

# Delta Ct (target - GAPDH), as requested in the sheet notes
ref = long[long.gene == "GAPDH"].set_index("sample_id")["ct"]
long["delta_ct_vs_gapdh"] = long["ct"] - long["sample_id"].map(ref)
long.loc[long.gene == "GAPDH", "delta_ct_vs_gapdh"] = np.nan
LOG.append("Hidden column J kept as run_id. The hard-coded summary table in the sheet was discarded "
           "and recomputed from the data (see below).")

long.to_csv(OUT, index=False, encoding="utf-8")

# ---------------------------------------------------------------- 8. recompute the summary
usable = long[(~long.excluded) & long.ct.notna() & long.gene.isin(["SEPT2", "MARCH1"])]
summary = (usable.groupby(["treatment", "dose_uM", "gene"])
                 .agg(mean_ct=("ct", "mean"), n=("ct", "size"))
                 .round(2).reset_index())

with open(LOG_OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(f"- {line}" for line in LOG))

if __name__ == "__main__":
    print(long.head(8).to_string(index=False))
    print("\nRows:", len(long), "| samples:", long.sample_id.nunique())
    print("\nFlags:\n", long.ct_flag.value_counts().to_string())
    print("\nRecomputed summary (excluded rows and non-numeric Ct left out):\n", summary.to_string(index=False))
